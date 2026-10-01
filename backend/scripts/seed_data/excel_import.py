from io import BytesIO
from time import perf_counter
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from scripts.seed_data.sme_retail import (
    adapt_sme_retail_row,
    is_sme_retail_schema,
)
from app.schemas.transaction import TransactionCreate


MAX_ROWS = 50_000


class ExcelImportError(ValueError):
    pass


def _normalize_cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def parse_transaction_workbook(content: bytes) -> dict[str, Any]:
    if not content:
        raise ExcelImportError("The uploaded Excel file is empty.")

    started_at = perf_counter()

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ExcelImportError(
            "The uploaded file is not a valid Excel workbook."
        ) from exc

    try:
        worksheet = workbook.active

        if worksheet is None:
            raise ExcelImportError(
                "The Excel workbook has no active worksheet."
            )

        rows = worksheet.iter_rows(values_only=True)

        try:
            raw_headers = list(next(rows))
        except StopIteration as exc:
            raise ExcelImportError(
                "The Excel worksheet is empty."
            ) from exc

        while raw_headers and raw_headers[-1] is None:
            raw_headers.pop()

        headers = [
            str(header).strip() if header is not None else ""
            for header in raw_headers
        ]

        if not headers or any(not header for header in headers):
            raise ExcelImportError(
                "Every used Excel column must have a header."
            )

        if len(headers) != len(set(headers)):
            raise ExcelImportError(
                "Excel column headers must be unique."
            )

        allowed_standard_headers = set(
            TransactionCreate.model_fields
        )

        is_standard_schema = (
            "id" in headers
            and set(headers).issubset(allowed_standard_headers)
        )

        is_retail_schema = is_sme_retail_schema(headers)

        if is_retail_schema:
            source_schema = "SME_RETAIL_EXPENSES"
        elif is_standard_schema:
            source_schema = "STANDARD_TRANSACTION"
        else:
            raise ExcelImportError(
                "The Excel column structure is not supported."
            )

        transactions: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []

        total_rows = 0
        complete_rows = 0
        partial_rows = 0
        rejected_rows = 0

        for excel_row_number, row in enumerate(rows, start=2):
            values = list(row[:len(headers)])

            if not any(value is not None for value in values):
                continue

            total_rows += 1

            if total_rows > MAX_ROWS:
                raise ExcelImportError(
                    f"The Excel file exceeds the "
                    f"{MAX_ROWS:,}-row limit."
                )

            source_row = {
                header: _normalize_cell(value)
                for header, value in zip(headers, values)
            }

            if is_retail_schema:
                transaction_input, metadata = (
                    adapt_sme_retail_row(source_row)
                )
            else:
                transaction_input = source_row
                metadata = {
                    "source_schema": "STANDARD_TRANSACTION"
                }

            try:
                transaction = TransactionCreate.model_validate(
                    transaction_input
                )
            except ValidationError as exc:
                rejected_rows += 1

                if len(errors) < 100:
                    messages = [
                        (
                            ".".join(
                                str(item)
                                for item in error["loc"]
                            )
                            + ": "
                            + error["msg"]
                        )
                        for error in exc.errors()
                    ]

                    errors.append(
                        {
                            "row": excel_row_number,
                            "messages": messages,
                        }
                    )

                continue

            transaction_data = transaction.model_dump(mode="json")

            missing_fields = [
                field_name
                for field_name, value in transaction_data.items()
                if value is None
            ]

            if missing_fields:
                partial_rows += 1
                data_quality_status = "PARTIAL"
            else:
                complete_rows += 1
                data_quality_status = "COMPLETE"

            transactions.append(
                {
                    **transaction_data,
                    "data_quality_status": data_quality_status,
                    "missing_fields": missing_fields,
                    "metadata": metadata,
                }
            )

        processing_time_seconds = (
            perf_counter() - started_at
        )

        return {
            "source_schema": source_schema,
            "total_rows": total_rows,
            "accepted_rows": len(transactions),
            "complete_rows": complete_rows,
            "partial_rows": partial_rows,
            "rejected_rows": rejected_rows,
            "processing_time_seconds": round(
                processing_time_seconds,
                4,
            ),
            "transactions": transactions,
            "errors": errors,
        }
    finally:
        workbook.close()