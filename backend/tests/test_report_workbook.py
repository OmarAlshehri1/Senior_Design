from io import BytesIO

from openpyxl import load_workbook

from app.services.report_workbook import build_report_workbook


def test_workbook_uses_persisted_values_and_contains_charts() -> None:
    report = {
        "id": "RPT-2026-10-03",
        "period_start": "2026-10-03T00:00:00+00:00",
        "period_end": "2026-10-04T00:00:00+00:00",
        "summary": {
            "daily_summary": {
                "total_transactions": 3,
                "transactions_evaluated": 3,
                "high_risk_transactions": 1,
                "average_risk_score": 45.5,
            },
            "risk_distribution": {"low": 1, "medium": 1, "high": 1},
            "alert_summary": {"total": 1, "active": 1, "reviewed": 0, "high": 1, "medium": 0},
            "rule_summary": [{"rule_key": "ghost_vendors", "rule_name": "Ghost Vendors", "violation_count": 1}],
            "high_risk_transactions": [{"id": "TX-HIGH-001", "amount": 500, "risk_score": 88}],
        },
    }
    workbook = load_workbook(BytesIO(build_report_workbook(report)))
    summary = workbook["Audit Summary"]
    details = workbook["High-Risk Transactions"]

    assert summary["A6"].value == 3
    assert summary["D6"].value == 3
    assert summary["G6"].value == 1
    assert summary["J6"].value == 45.5
    assert [summary[f"B{row}"].value for row in (12, 13, 14)] == [1, 1, 1]
    assert summary["A31"].value == "Ghost Vendors"
    assert summary["B31"].value == 1
    assert len(summary._charts) == 2
    assert details["A5"].value == "TX-HIGH-001"
    assert details["D5"].value == 500
    assert details["F5"].value == 88


def test_workbook_handles_empty_report_without_fake_chart_data() -> None:
    content = build_report_workbook({"id": "RPT-EMPTY", "summary": {}})
    sheet = load_workbook(BytesIO(content))["Audit Summary"]
    assert sheet["A6"].value == "N/A"
    assert sheet["D13"].value == "No transactions in this reporting period."
    assert sheet["A31"].value == "No rule violations recorded."
    assert not sheet._charts


def test_vendor_name_is_kept_as_text() -> None:
    content = build_report_workbook({
        "id": "RPT-2026-10-03",
        "summary": {
            "high_risk_transactions": [
                {"id": "TX-1", "vendor_name": "=HYPERLINK(\"https://example.com\")"}
            ]
        },
    })
    cell = load_workbook(BytesIO(content))["High-Risk Transactions"]["C5"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")
