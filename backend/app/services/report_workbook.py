from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "172536"
BLUE = "0876B9"
PALE = "E7F2FA"
LIGHT = "F4F7FA"
WHITE = "FFFFFF"


def _number(value: Any) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value) if "." in str(value) else int(value)
    except (TypeError, ValueError):
        return None


def _heading(sheet: Any, row: int, start: int, end: int, label: str) -> None:
    sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    cell = sheet.cell(row, start, label)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos", size=12, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 26
    for col in range(start + 1, end + 1):
        sheet.cell(row, col).fill = PatternFill("solid", fgColor=NAVY)


def _table_header(sheet: Any, row: int, labels: list[str], start: int = 1) -> None:
    for col, label in enumerate(labels, start):
        cell = sheet.cell(row, col, label)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 29


def _table_row(sheet: Any, row: int, values: list[Any], start: int = 1) -> None:
    for col, value in enumerate(values, start):
        cell = sheet.cell(row, col, value)
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"
        cell.fill = PatternFill("solid", fgColor=PALE if row % 2 == 0 else WHITE)
        cell.font = Font(name="Aptos", size=10, color=NAVY)
        cell.border = Border(bottom=Side(style="hair", color="CFDCE7"))
        cell.alignment = Alignment(vertical="center")
        if isinstance(value, (int, float)):
            cell.number_format = "#,##0.##"
    sheet.row_dimensions[row].height = 23


def _metric(sheet: Any, row: int, col: int, label: str, value: Any) -> None:
    sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
    sheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 2)
    label_cell = sheet.cell(row, col, label)
    label_cell.font = Font(name="Aptos", size=10, color="49627A")
    value_cell = sheet.cell(row + 1, col, value if value is not None else "N/A")
    value_cell.font = Font(name="Aptos Display", size=23, bold=True, color=NAVY)
    value_cell.number_format = "#,##0.##"
    for r in range(row, row + 3):
        for c in range(col, col + 3):
            sheet.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT)
    sheet.row_dimensions[row + 1].height = 25
    sheet.row_dimensions[row + 2].height = 16


def build_report_workbook(report: dict[str, Any]) -> bytes:
    """Render a completed report's persisted summary as a formatted Excel file."""
    summary = report.get("summary")
    if not isinstance(summary, dict):
        raise ValueError("Completed report summary is unavailable.")

    daily = summary.get("daily_summary") or {}
    risks = summary.get("risk_distribution") or {}
    alerts = summary.get("alert_summary") or {}
    rules = summary.get("rule_summary") or []
    transactions = summary.get("high_risk_transactions") or []

    book = Workbook()
    sheet = book.active
    sheet.title = "Audit Summary"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"
    for col in range(1, 13):
        sheet.column_dimensions[get_column_letter(col)].width = 13
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["D"].width = 18

    sheet.merge_cells("A1:L2")
    title = sheet["A1"]
    title.value = "Continuous Auditing | Daily Report"
    title.font = Font(name="Aptos Display", size=23, bold=True, color=BLUE)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 29
    sheet.row_dimensions[2].height = 20
    sheet.merge_cells("A3:L3")
    sheet["A3"] = (
        f"Report {report.get('id', '')}  |  "
        f"UTC period: {report.get('period_start', '')} to {report.get('period_end', '')} (exclusive)"
    )
    sheet["A3"].font = Font(name="Aptos", size=10, color="49627A")

    _metric(sheet, 5, 1, "Total transactions", _number(daily.get("total_transactions")))
    _metric(sheet, 5, 4, "Transactions evaluated", _number(daily.get("transactions_evaluated")))
    _metric(sheet, 5, 7, "High-risk transactions", _number(daily.get("high_risk_transactions")))
    _metric(sheet, 5, 10, "Average risk score / 100", _number(daily.get("average_risk_score")))

    _heading(sheet, 10, 1, 12, "Risk distribution")
    _table_header(sheet, 11, ["Risk level", "Transactions"])
    risk_labels = [("Low", "low"), ("Medium", "medium"), ("High", "high")]
    for row, (label, key) in enumerate(risk_labels, 12):
        _table_row(sheet, row, [label, _number(risks.get(key)) or 0])
    risk_total = sum(sheet.cell(row, 2).value for row in range(12, 15))
    if risk_total:
        chart = PieChart()
        chart.title = "Transactions by risk level"
        chart.add_data(Reference(sheet, min_col=2, min_row=11, max_row=14), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=12, max_row=14))
        chart.height = 8
        chart.width = 18
        sheet.add_chart(chart, "D12")
    else:
        sheet["D13"] = "No transactions in this reporting period."

    _heading(sheet, 29, 1, 12, "Audit rule violations")
    _table_header(sheet, 30, ["Rule", "Violations"])
    valid_rules = [rule for rule in rules if isinstance(rule, dict)]
    for row, rule in enumerate(valid_rules, 31):
        _table_row(sheet, row, [rule.get("rule_name") or rule.get("rule_key") or "Unknown", _number(rule.get("violation_count")) or 0])
    if valid_rules:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Violations by audit rule"
        chart.y_axis.title = "Violations"
        chart.add_data(Reference(sheet, min_col=2, min_row=30, max_row=30 + len(valid_rules)), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=31, max_row=30 + len(valid_rules)))
        chart.height = 8
        chart.width = 18
        sheet.add_chart(chart, "D31")
    else:
        sheet["A31"] = "No rule violations recorded."

    alert_row = max(50, 33 + len(valid_rules))
    _heading(sheet, alert_row, 1, 12, "Alert summary")
    _table_header(sheet, alert_row + 1, ["Alert status", "Count"])
    for offset, (label, key) in enumerate([
        ("Total", "total"), ("Active", "active"), ("Reviewed", "reviewed"),
        ("High risk", "high"), ("Medium risk", "medium"),
    ], 2):
        _table_row(sheet, alert_row + offset, [label, _number(alerts.get(key)) or 0])

    details = book.create_sheet("High-Risk Transactions")
    details.sheet_view.showGridLines = False
    details.freeze_panes = "A5"
    details.merge_cells("A1:G2")
    details["A1"] = "High-Risk Transactions"
    details["A1"].font = Font(name="Aptos Display", size=22, bold=True, color=BLUE)
    details["A1"].alignment = Alignment(vertical="center")
    details.merge_cells("A3:G3")
    details["A3"] = "Transactions included in the persisted daily report."
    details["A3"].font = Font(name="Aptos", size=10, color="49627A")
    _table_header(details, 4, ["Transaction ID", "Timestamp (UTC)", "Vendor", "Amount", "Currency", "Risk score", "Risk level"])
    for row, transaction in enumerate(transactions, 5):
        if not isinstance(transaction, dict):
            continue
        _table_row(details, row, [
            transaction.get("id"), transaction.get("timestamp"), transaction.get("vendor_name"),
            _number(transaction.get("amount")), transaction.get("currency"),
            _number(transaction.get("risk_score")), transaction.get("risk_level"),
        ])
        details.cell(row, 4).number_format = "#,##0.00"
    details.auto_filter.ref = f"A4:G{max(4, details.max_row)}"
    for col, width in {"A": 25, "B": 28, "C": 30, "D": 16, "E": 12, "F": 15, "G": 14}.items():
        details.column_dimensions[col].width = width

    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
