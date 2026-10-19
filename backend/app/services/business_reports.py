"""Memoranda from saved evidence, with separately labelled optional AI commentary."""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from html import escape
import json
from typing import Any

from app.repositories import identity
from app.services.audit_reports import normalize_daily_period
from app.services.report_narratives import enrich_findings

RULES = {
    'segregation_of_duties': ('Segregation of Duties', 'Verify creator/approver separation and supporting approval evidence.'),
    'approval_limits': ('Approval Limits', 'Check the recorded authorization limit and supporting approval.'),
    'duplicate_payment': ('Duplicate Payments', 'Compare invoices and payment history before deciding whether a duplicate occurred.'),
    'invoice_splitting': ('Invoice Splitting', 'Review related invoices within the rule window and their approval evidence.'),
    'ghost_vendors': ('Ghost Vendors', 'Reconcile the vendor against the approved registry and validate registration status.'),
}

GLOSSARY = ('A failed-rule occurrence is one failed rule for one transaction; several may refer to the same transaction. '
            'Coverage counts transactions with at least one executed rule. A reviewed alert is not necessarily remediated. '
            'Transaction value is not loss. The extraction date differs from the transaction period. Unscored is not low risk.')


def _money_totals(rows: list[dict]) -> list[dict]:
    totals: dict[str, Decimal] = defaultdict(Decimal)
    counts: Counter = Counter()
    for row in rows:
        if row.get('amount') is not None and row.get('currency'):
            currency = row['currency']
            totals[currency] += Decimal(str(row['amount']))
            counts[currency] += 1
    return [{'currency': key, 'amount': str(value), 'transactions': counts[key]}
            for key, value in sorted(totals.items())]


def build_document(source: dict, *, recipient: str, sender: str, data_context: str) -> dict:
    transactions = source['transactions']
    alerts, cases = source['alerts'], source['cases']
    total = len(transactions)
    evaluated = sum(any(r.get('status') in ('PASSED', 'FAILED') for r in t['rule_results']) for t in transactions)
    risks = Counter(t.get('risk_level') for t in transactions)
    scores = [Decimal(str(t['risk_score'])) for t in transactions if t.get('risk_score') is not None]
    high = [t for t in transactions if t.get('risk_level') == 'HIGH']
    affected = {t['id'] for t in transactions if any(r.get('status') == 'FAILED' for r in t['rule_results'])}
    rules = []
    recommendations = []
    findings = [{
        'id': 'F1', 'title': 'Risk priorities',
        'text': f"{len(high)} of {total} transactions are classified HIGH; {risks['MEDIUM']} are MEDIUM and {risks['LOW']} LOW. "
                f"{total-len(scores)} have no stored combined risk score. Recorded transaction values are not confirmed losses or validated financial exposure.",
        'transaction_ids': [t['id'] for t in high],
    }]
    if high:
        recommendations.append({'finding_id': 'F1', 'priority': 'HIGH', 'proposed_owner': 'Audit Supervisor / assigned Auditor',
                                'action': 'Review high-risk records and supporting documents; document each investigation outcome.'})
    for key, (name, recommendation) in RULES.items():
        executed, failed, reasons = set(), set(), Counter()
        for transaction in transactions:
            results = [r for r in transaction['rule_results'] if r.get('rule_key') == key]
            if any(r.get('status') in ('PASSED', 'FAILED') for r in results):
                executed.add(transaction['id'])
            else:
                evidence = (results[0].get('evidence') or {}) if results else {}
                reason = ('SNAPSHOT_OR_RULE_MISSING' if not results else
                          'MISSING_REQUIRED_FIELDS' if evidence.get('missing_fields') or evidence.get('missing_any_of') else
                          'CONTEXT_UNAVAILABLE' if evidence.get('historical_context_available') is False or evidence.get('vendor_registry_available') is False else
                          'NOT_EVALUATED')
                reasons[reason] += 1
            if any(r.get('status') == 'FAILED' for r in results):
                failed.add(transaction['id'])
        rules.append({'key': key, 'name': name, 'evaluated': len(executed), 'not_evaluated': total-len(executed),
                      'failed': len(failed), 'transaction_ids': sorted(failed), 'exclusions': dict(reasons)})
        if failed:
            reference = f'F{len(findings)+1}'
            findings.append({'id': reference, 'title': name,
                             'text': f"{len(failed)} transactions have a stored FAILED result for {name}, out of {len(executed)} evaluated for this rule. "
                                     'These are indicators for review, not confirmed fraud or proof of loss. Root cause is not established.',
                             'transaction_ids': sorted(failed)})
            recommendations.append({'finding_id': reference, 'priority': 'HIGH' if any(t['id'] in failed for t in high) else 'MEDIUM',
                                    'proposed_owner': 'Audit Supervisor / relevant control owner', 'action': recommendation})
    followup_ref = f'F{len(findings)+1}'
    active_alerts = sum(a['status'] == 'ACTIVE' for a in alerts)
    findings.append({'id': followup_ref, 'title': 'Follow-up at the evidence extraction date',
                     'text': f"{len(alerts)} authorized related alerts: {active_alerts} ACTIVE and {sum(a['status']=='REVIEWED' for a in alerts)} REVIEWED. "
                             f"{len(cases)} authorized related cases, including {sum(c['status']=='CLOSED' for c in cases)} CLOSED. "
                             'These are current states at the extraction date, not reconstructed period-end states. Reviewed does not mean remediated.',
                     'transaction_ids': []})
    if active_alerts:
        recommendations.append({'finding_id': followup_ref, 'priority': 'HIGH', 'proposed_owner': 'Audit Supervisor',
                                'action': 'Assign and investigate outstanding alerts; record decisions and linked cases where appropriate.'})
    coverage_ref = f'F{len(findings)+1}'
    findings.append({'id': coverage_ref, 'title': 'Evaluation coverage and data limitations',
                     'text': f'{evaluated} of {total} transactions have at least one executed rule. '
                             f"{sum(t.get('data_quality_status')=='PARTIAL' for t in transactions)} are marked PARTIAL for data quality. "
                             'This does not imply all five rules executed. Per-rule eligibility and exclusion reasons are in Appendix B.',
                     'transaction_ids': [t['id'] for t in transactions if not any(r.get('status') in ('PASSED','FAILED') for r in t['rule_results'])]})
    if evaluated < total or any(r['not_evaluated'] for r in rules):
        recommendations.append({'finding_id': coverage_ref, 'priority': 'MEDIUM', 'proposed_owner': 'Audit Supervisor / data owner',
                                'action': 'Review eligibility and missing context before making broader control-coverage claims.'})
    scope_description = ('All dated organization transactions in the selected UTC day.' if source['scope'] == 'ORGANIZATION' else
                         'Only dated transactions linked to currently assigned team alerts or team cases. Unlinked transactions and other teams are excluded.')
    coverage = str((Decimal(evaluated)*100/total).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) if total else None
    context = {'DEMONSTRATION': 'User-declared demonstration data; findings are for workflow validation, not real business conclusions.',
               'OPERATIONAL': 'User-declared operational data; this classification is not independent source validation.',
               'UNVERIFIED': 'Data provenance is unverified; synthetic and operational records have not been independently distinguished.'}[data_context]
    document = {
        'schema_version': '2.0.0', 'approval_status': 'DRAFT', 'generation_status': 'COMPLETED',
        'title': 'Daily Transaction Risk and Control Exceptions — Business Memorandum',
        'organization_name': source['organization_name'], 'recipient': recipient, 'sender': sender,
        'prepared_by': deepcopy(source['prepared_by']), 'data_context': data_context,
        'scope': source['scope'], 'team_id': source.get('team_id'), 'team_name': source.get('team_name'),
        'period_start': source['period_start'], 'period_end': source['period_end'], 'as_of': source['as_of'],
        'review': {'reviewer': None, 'designation': None, 'reviewed_at': None, 'signature': None},
        'terms': f"This draft helps management prioritize investigation and control review. {scope_description} "
                 'It is not an external audit opinion, assurance of complete control effectiveness, or confirmation of fraud. '
                 'Direct ERP/POS integration is not established by this report.',
        'methodology': 'Transactions are selected by their transaction timestamp within a half-open UTC day. '
                       'The latest persisted evaluation and combined-risk snapshots at extraction are used without re-running rules or the model. '
                       'The five rules are Segregation of Duties, Approval Limits, Duplicate Payments, Invoice Splitting and Ghost Vendors. '
                       'Isolation Forest supplies the anomaly component. Risk = 0.60 × rule score + 0.40 × anomaly score. '
                       'LOW < 50; MEDIUM ≥ 50 and < 75; HIGH ≥ 75. Scores are review-priority indicators, not fraud probabilities. '
                       'Evaluation coverage means at least one rule executed, not all five. Per-rule denominators are shown separately. '
                       'Rule evidence and risk inputs have their own source IDs, versions and dates in the appendices; they may be from different evaluations. '
                       'Alert/case states are captured at extraction and may post-date the transaction period.',
        'limitations': [context, f"{source.get('undated_in_scope', 0)} undated transactions in the authorized cohort are excluded from date-bounded reports; this count is not a count for the selected day.",
                        'No root cause, monetary loss, policy breach adjudication, or remediation is inferred from a score or failed rule.',
                        'Currency totals include only records with both amount and currency. Missing values are not assumed to be zero.',
                        'No historical period-end workflow state is reconstructed. No external supporting documents were examined by the generator.'],
        'metrics': {'total_transactions': total, 'evaluated_transactions': evaluated, 'coverage_percent': coverage,
                    'scored_transactions': len(scores), 'average_risk_score': str((sum(scores)/len(scores)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) if scores else None,
                    'risk_distribution': {key: risks[key] for key in ('LOW', 'MEDIUM', 'HIGH')},
                    'unscored_transactions': total-len(scores), 'unique_exception_transactions': len(affected),
                    'rule_failure_occurrences': sum(r['failed'] for r in rules),
                    'amount_totals': _money_totals(transactions), 'high_risk_amount_totals': _money_totals(high),
                    'missing_amount_or_currency': sum(t.get('amount') is None or not t.get('currency') for t in transactions)},
        'rules': rules, 'findings': findings, 'recommendations': recommendations,
        'conclusion': ('No dated transactions fall within this authorized scope and period. No assurance conclusion can be drawn.' if not total else
                       f"The {len(affected)} transactions with rule exceptions and {len(high)} high-risk classifications identify review priorities. "
                       'These groups may overlap and must not be added together. Results do not establish fraud, loss, or overall control effectiveness. '
                       'Interpret the conclusions within the stated data context and eligibility limitations.'),
        'evidence': deepcopy({'transactions': transactions, 'alerts': alerts, 'cases': cases}),
    }
    # Always provide detailed explanations, even offline or without a provider.
    return enrich_findings(document, requested=False)


def generate_business_report(*, actor_id: str, period_start=None, period_end=None,
                             recipient='Management and Audit Supervisor', sender='Continuous Auditing System — automated preparation',
                             data_context='UNVERIFIED', use_gemini=True) -> dict:
    if period_start is not None:
        period_start, _ = normalize_daily_period(period_start, period_end)
    source = identity.rpc('get_business_report_source', p_actor=actor_id,
                          p_start=period_start.isoformat() if period_start else None)
    if not isinstance(source, dict) or not all(k in source for k in ('transactions','alerts','cases','organization_name','as_of')):
        raise identity.IdentityError('REPORT_SOURCE_UNAVAILABLE')
    document = build_document(source, recipient=recipient, sender=sender, data_context=data_context)
    document = enrich_findings(document, requested=use_gemini)
    saved = identity.rpc('save_business_report', p_actor=actor_id, p_document=document)
    if not isinstance(saved, dict) or not saved.get('id'):
        raise identity.IdentityError('REPORT_SAVE_UNAVAILABLE')
    return saved


def read_business_reports(*, actor_id: str, report_id=None, page=1, page_size=25):
    return identity.rpc('read_business_reports', p_actor=actor_id, p_id=str(report_id) if report_id else None,
                        p_page=page, p_page_size=page_size)


def render_business_html(report: dict) -> str:
    d = report['document']
    def text(value):
        return escape('Not recorded' if value is None else str(value), quote=True)
    def table(headers, rows):
        return '<table><thead><tr>' + ''.join(f'<th>{text(h)}</th>' for h in headers) + '</tr></thead><tbody>' + ''.join(
            '<tr>' + ''.join(f'<td>{text(v)}</td>' for v in row) + '</tr>' for row in rows) + '</tbody></table>'
    def section(title, body):
        return f'<section><h2>{text(title)}</h2>{body}</section>'
    def paragraph(value):
        return f'<p>{text(value)}</p>'
    def reference(*values):
        return ' / '.join('Not recorded' if value is None else str(value) for value in values)
    m = d['metrics']
    content = '<h1>Memorandum</h1><p class="draft">AUTOMATICALLY GENERATED DRAFT — HUMAN REVIEW REQUIRED</p>'
    header = [
        ['Organization',d['organization_name']], ['To',d['recipient']], ['From',d['sender']],
        ['Subject',d['title']], ['Report / revision ID',f"BR-{report['id']}"],
        ['Prepared at',report['created_at']], ['Evidence as of',d['as_of']],
        ['Transaction period (UTC)',f"{d['period_start']} to {d['period_end']} (exclusive)"],
        ['Scope',d['scope'] + (f" — {d['team_name']}" if d.get('team_name') else '')], ['Data context',d['data_context']],
    ]
    content += '<dl class="memo-header">' + ''.join(f'<dt>{text(k)}</dt><dd>{text(v)}</dd>' for k,v in header) + '</dl>'
    content += section('I. Terms of Reference',paragraph(d['terms']))
    content += section('II. Procedure / Methodology',paragraph(d['methodology']) + '<ul>' + ''.join(f'<li>{text(x)}</li>' for x in d['limitations']) + '</ul>')
    findings = table(['Metric','Recorded value'], [
        ['Transactions',m['total_transactions']], ['Evaluated (at least one rule)',m['evaluated_transactions']],
        ['Evaluation coverage (%)',m['coverage_percent']], ['Stored combined scores',m['scored_transactions']],
        ['Average risk score / 100',m['average_risk_score']], ['Unscored',m['unscored_transactions']],
        ['Unique transactions with rule failures',m['unique_exception_transactions']], ['Failed-rule occurrences (may overlap)',m['rule_failure_occurrences']],
        ['Missing amount or currency',m['missing_amount_or_currency']],
    ])
    narrative = d.get('narrative') or {}
    findings += paragraph('Finding explanations: ' + ('AI-assisted (Gemini); advisory, human review required.'
                          if narrative.get('provider')=='GEMINI' else 'Standard explanations; AI commentary was not used.'))
    for finding in d['findings']:
        findings += f"<h3>{text(finding['id'])}. {text(finding['title'])}</h3>" + paragraph(finding['text'])
        findings += paragraph('Evidence: ' + (', '.join(finding['transaction_ids']) or 'Authorized workflow schedules in Appendix C.'))
        analysis = finding.get('analysis') or {}
        for key,label in [('interpretation','Interpretation'),('business_implication','Potential business implication'),
                          ('recommended_review','Proposed verification'),('limitations','Limits of interpretation')]:
            if analysis.get(key):
                findings += f'<h4>{text(label)}</h4>' + paragraph(analysis[key])
    findings += table(['Currency','Recorded transaction value','Records with amount and currency'],
                      [[x['currency'],x['amount'],x['transactions']] for x in m['amount_totals']])
    findings += paragraph('High-risk recorded values only (not losses):') + table(['Currency','Amount','Transactions'],
                      [[x['currency'],x['amount'],x['transactions']] for x in m['high_risk_amount_totals']])
    content += section('III. Findings',findings)
    content += section('IV. Conclusion',paragraph(d['conclusion']))
    content += section('V. Recommendations',paragraph('Proposed actions only; no approval, due date or completed remediation is implied.') +
                       table(['Finding','Priority','Proposed responsible role','Proposed action'],
                             [[r['finding_id'],r['priority'],r['proposed_owner'],r['action']] for r in d['recommendations']]))
    content += section('VI. Review / Sign-off',table(['Field','Value'],[[k,v] for k,v in d['review'].items()]) + paragraph('Draft — not approved. Generation completion is not review or sign-off.'))
    content += section('Appendix A — Transaction schedule',table(
        ['Transaction','UTC timestamp','Vendor','Amount','Currency','Data quality','Missing fields','Risk score','Risk level','Risk snapshot / version / date','Rule / anomaly versions used in risk'],
        [[t['id'],t['timestamp'],t.get('vendor_name'),t.get('amount'),t.get('currency'),t.get('data_quality_status'),
          ', '.join(t.get('missing_fields') or []),t.get('risk_score'),t.get('risk_level'),
          reference(t.get('risk_snapshot_id'),t.get('scoring_version'),t.get('calculated_at')),
          reference(t.get('rule_evaluation_version'),t.get('anomaly_model_version')) ] for t in d['evidence']['transactions']]))
    rules = table(['Rule','Evaluated','Not evaluated','Failed transactions','Exclusion reasons'],
                  [[r['name'],r['evaluated'],r['not_evaluated'],r['failed'],json.dumps(r['exclusions'],sort_keys=True)] for r in d['rules']])
    rules += table(['Transaction','Evaluation ID / version / date','Rule / version','Status','Recorded detail','Recorded evidence'],[
        [t['id'],reference(t.get('evaluation_id'),t.get('evaluation_version'),t.get('evaluated_at')),
         reference(r.get('rule_name'),r.get('rule_version')),r.get('status'),r.get('detail'),json.dumps(r.get('evidence') or {},sort_keys=True)]
        for t in d['evidence']['transactions'] for r in t['rule_results']])
    risks = table(['Classification','Transactions'],[[key,value] for key,value in m['risk_distribution'].items()])
    risks += '<div class="bars" role="img" aria-label="Risk distribution; exact counts in the table above">' + ''.join(
        f'<div>{text(key)}: {value}<span style="width:{round(value*100/max(m["total_transactions"],1),2)}%"></span></div>'
        for key,value in m['risk_distribution'].items()) + '</div>'
    content += section('Appendix B — Coverage, rule evidence and risk distribution',rules + risks)
    content += section('Appendix C — Authorized follow-up at extraction', paragraph('Snapshot as of ' + d['as_of']) +
        table(['Alert','Transaction','Severity','Status','Created','Reviewed'],[
            [a['id'],a['transaction_id'],a['severity'],a['status'],a.get('created_at'),a.get('reviewed_at')] for a in d['evidence']['alerts']]) +
        table(['Case','Transaction','Status','Priority','SLA','Created','Updated','Closed','Outcome'],[
            [f"CASE-{c['case_number']:06}",c['transaction_id'],c['status'],c['priority'],c['sla_status'],c['created_at'],c['updated_at'],c.get('closed_at'),c.get('resolution_outcome')]
            for c in d['evidence']['cases']]))
    content += section('Appendix D — Glossary',paragraph(GLOSSARY) +
        table(['Document metadata','Saved value'],[['Schema version',d['schema_version']],['Requested by',d['prepared_by'].get('name')],
            ['Preparation','Automated reporting service'],['Generation status',d['generation_status']],['Approval status',d['approval_status']],
            ['Commentary provider',narrative.get('provider')],['Commentary status',narrative.get('status')],
            ['Commentary model',narrative.get('model')],['Commentary prompt version',narrative.get('prompt_version')],
            ['Commentary attempted model',narrative.get('attempted_model')],['Commentary failure code',narrative.get('failure_code')],
            ['Commentary attempts',narrative.get('attempts')],
            ['Commentary generated at',narrative.get('generated_at')]]))
    return '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">' + \
        '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">' + \
        '<title>' + text(d['title']) + '</title><style>' + \
        'body{font:15px/1.6 Georgia,serif;color:#172536;max-width:1100px;margin:32px auto;padding:24px;background:white}' + \
        'h1{font-size:34px}h2{border-bottom:2px solid #0876b9;padding-bottom:8px}h3{margin-bottom:6px}' + \
        '.memo-header{display:grid;grid-template-columns:180px 1fr;gap:8px 18px}.memo-header dt{font-weight:bold}.memo-header dd{margin:0;overflow-wrap:anywhere}' + \
        'table{border-collapse:collapse;width:100%;margin:18px 0;font:12px/1.5 Arial,sans-serif;overflow-wrap:anywhere}' + \
        'th,td{padding:9px;border:1px solid #cbd5e1;text-align:left;vertical-align:top}th{background:#edf3f8}' + \
        '.draft{border:1px solid #b45309;padding:12px}.bars span{display:block;height:12px;background:#0876b9}' + \
        '@media print{@page{size:A4;margin:14mm}body{margin:0;padding:0;font-size:11pt}thead{display:table-header-group}tr{break-inside:avoid}h2,h3{break-after:avoid}table{font-size:8pt}}' + \
        '</style></head><body>' + content + '</body></html>'


def build_business_workbook(report: dict) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment
    from app.services.report_workbook import _table_header, _table_row
    d = report['document']
    book = Workbook()
    sheets = [
        ('Memorandum', ['Section','Content'], [
            ['Report ID',f"BR-{report['id']}"], ['Organization',d['organization_name']], ['To',d['recipient']], ['From',d['sender']],
            ['Generated / saved at',report['created_at']], ['Document version',d['schema_version']], ['Requested by',d['prepared_by'].get('name')],
            ['Period',f"{d['period_start']} to {d['period_end']} (exclusive)"], ['Evidence as of',d['as_of']],
            ['Status','DRAFT — human review required'], ['Data context',d['data_context']], ['Scope',d['scope']],
            ['I. Terms of Reference',d['terms']], ['II. Methodology',d['methodology']],
            *[['Limitation',x] for x in d['limitations']],
            *[[f"III. {f['id']} — {f['title']}",f['text'] + ' Evidence: ' + (', '.join(f['transaction_ids']) or 'Authorized follow-up schedules.')] for f in d['findings']],
            *[[f"III. {f['id']} — {key.replace('_',' ').title()} ({f.get('analysis_source','BASELINE')})",value]
              for f in d['findings'] for key,value in (f.get('analysis') or {}).items()],
            ['IV. Conclusion',d['conclusion']],
            *[[f"V. {r['finding_id']} — {r['priority']}",r['proposed_owner'] + ': ' + r['action']] for r in d['recommendations']],
            *[[f'VI. {k}',v] for k,v in d['review'].items()],
            ['Glossary',GLOSSARY],
            *[[f'Commentary {key}',value] for key,value in (d.get('narrative') or {}).items()],
        ]),
        ('Metrics', ['Metric','Value'], [[k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v] for k,v in d['metrics'].items()]),
        ('Transactions', ['ID','Timestamp UTC','Vendor','Amount','Currency','Data quality','Missing fields','Risk score','Risk level','Risk snapshot','Risk version','Calculated at','Rule score used','AI score used'],
         [[t.get(k) if k!='missing_fields' else ', '.join(t.get(k) or []) for k in ('id','timestamp','vendor_name','amount','currency','data_quality_status','missing_fields','risk_score','risk_level','risk_snapshot_id','scoring_version','calculated_at','rule_score','ai_score')] for t in d['evidence']['transactions']]),
        ('Rule Coverage', ['Rule','Evaluated','Not evaluated','Failed','Exclusions'],
         [[r['name'],r['evaluated'],r['not_evaluated'],r['failed'],json.dumps(r['exclusions'])] for r in d['rules']]),
        ('Rule Evidence', ['Transaction','Evaluation ID','Evaluation version','Evaluated at','Rule','Rule version','Status','Detail','Evidence'],
         [[t['id'],t.get('evaluation_id'),t.get('evaluation_version'),t.get('evaluated_at'),r.get('rule_name'),r.get('rule_version'),r.get('status'),r.get('detail'),json.dumps(r.get('evidence') or {})] for t in d['evidence']['transactions'] for r in t['rule_results']]),
        ('Alerts', ['ID','Transaction','Severity','Status','Created','Reviewed'],
         [[a.get(k) for k in ('id','transaction_id','severity','status','created_at','reviewed_at')] for a in d['evidence']['alerts']]),
        ('Cases', ['Number','Transaction','Status','Priority','SLA','Created','Updated','Closed','Outcome'],
         [[c.get(k) for k in ('case_number','transaction_id','status','priority','sla_status','created_at','updated_at','closed_at','resolution_outcome')] for c in d['evidence']['cases']]),
        ('Risk Distribution', ['Risk level','Transactions'], [[k,v] for k,v in d['metrics']['risk_distribution'].items()]),
    ]
    for index,(name,headers,rows) in enumerate(sheets):
        sheet = book.active if index==0 else book.create_sheet()
        sheet.title = name
        sheet.freeze_panes = 'A2'
        _table_header(sheet,1,headers)
        for row,values in enumerate(rows,2):
            _table_row(sheet,row,[v if v is not None else 'Not recorded' for v in values])
            for cell in sheet[row]:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            sheet.row_dimensions[row].height = min(409, max(40, 16 * (len(str(values[1] or '')) // 75 + 1))) if name=='Memorandum' else 38
        from openpyxl.utils import get_column_letter
        for col in range(1,len(headers)+1):
            sheet.column_dimensions[get_column_letter(col)].width = 90 if name=='Memorandum' and col==2 else 28
        sheet.auto_filter.ref = sheet.dimensions
    if d['metrics']['total_transactions']:
        sheet = book['Risk Distribution']
        chart = BarChart()
        chart.title = 'Recorded risk classifications'
        chart.add_data(Reference(sheet,min_col=2,min_row=1,max_row=4),titles_from_data=True)
        chart.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=4))
        sheet.add_chart(chart,'D2')
    output = BytesIO()
    book.save(output)
    return output.getvalue()
